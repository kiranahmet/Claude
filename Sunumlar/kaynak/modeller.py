# Davranış değişikliği modelleri: soru-cevap hazırlık notu (Word)
import sys
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SABLON, CIKIS = sys.argv[1], sys.argv[2]
KOYU, VURGU, GRI = RGBColor(0x0F, 0x4C, 0x5C), RGBColor(0x2A, 0x9D, 0x8F), RGBColor(0x5F, 0x71, 0x78)

d = docx.Document(SABLON)
body = d.element.body
for el in list(body):
    if not el.tag.endswith("}sectPr"):
        body.remove(el)

def par(text="", style=None, size=None, bold=False, color=None, italic=False, after=6):
    p = d.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    if text:
        add(p, text, size=size, bold=bold, color=color, italic=italic)
    return p

def add(p, text, size=None, bold=False, color=None, italic=False):
    # **kalın** işaretli parçaları kalın yazar
    parts = text.split("**")
    for i, t in enumerate(parts):
        if not t:
            continue
        r = p.add_run(t)
        if size: r.font.size = Pt(size)
        r.bold = (bold or i % 2 == 1) or None
        r.italic = italic or None
        if color: r.font.color.rgb = color

def madde(text, level=0):
    p = par(after=3)
    p.paragraph_format.left_indent = Cm(0.6 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    add(p, ("• " if level == 0 else "– ") + text)
    return p

def numara(n, text):
    p = par(after=3)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    add(p, f"{n}. " + text)

def etiket(text):
    par(text, bold=True, color=KOYU, after=3)

def golge(cell, renk):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), renk)
    tcPr.append(shd)

def tablo(basliklar, satirlar, genislik):
    t = d.add_table(rows=1, cols=len(basliklar))
    tblPr = t._tbl.tblPr
    kb = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "C9DCDF")
        kb.append(e)
    tblPr.append(kb)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, b in enumerate(basliklar):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(b); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        golge(c, "0F4C5C")
    for s in satirlar:
        cells = t.add_row().cells
        for i, v in enumerate(s):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(10)
    for row in t.rows:
        for i, w in enumerate(genislik):
            row.cells[i].width = Cm(w)
    par(after=4)

def not_kutusu(text):
    p = par(after=8)
    p.paragraph_format.left_indent = Cm(0.4)
    add(p, text, italic=True, color=GRI, size=10)

# ---------------- Başlık ----------------
par("Davranış Değişikliği Modelleri", size=22, bold=True, color=KOYU, after=2)
par("Sunumun 5. slaytı için soru-cevap hazırlık notu", size=13, color=VURGU, after=4)
par("Makalede bu beş model yalnızca adıyla sayılmaktadır (Giriş, kaynak 10–11). Hocanın ders slaytlarında bunlardan yalnızca Transteorik Model yer almaktadır; aşama adları 50. slayttaki gibi yazılmıştır. Diğer dört model genel literatür bilgisidir. “Bağlantı” notları hocanın slaytında veya makalede yazmayan, benim kurduğum benzetmelerdir.", size=10, color=GRI, after=10)

# ---------------- 1 ----------------
par("1. Sağlık İnanç Modeli", style="Heading 1")
par("**Kim, ne zaman:** 1950'lerde ABD Halk Sağlığı Servisi'nde sosyal psikologlar tarafından geliştirilmiştir (Rosenstock ve ark.).")
par("**Ana fikir:** Kişi bir hastalığın kendisi için gerçek ve ciddi bir tehdit olduğuna inanırsa ve davranışı değiştirmenin yararını engellerinden büyük görürse davranışını değiştirir.")
etiket("Bileşenleri ve sigara örneği")
tablo(["Bileşen", "Anlamı", "Sigara örneği"], [
    ["Algılanan duyarlılık", "“Bu hastalık bana da olabilir.”", "“Ben de KOAH ya da kanser olabilirim.”"],
    ["Algılanan ciddiyet", "“Olursa sonuçları ağır olur.”", "“Akciğer kanseri öldürücüdür.”"],
    ["Algılanan yarar", "“Değiştirirsem kazancım olur.”", "“Bırakırsam nefesim açılır, param kalır.”"],
    ["Algılanan engel", "Değiştirmenin zorlukları", "“Kilo alırım, sinirlenirim.”"],
    ["Harekete geçirici ipucu", "Kişiyi o an harekete geçiren olay", "Hekim uyarısı, yakınının hastalanması"],
    ["Öz-etkililik (sonradan eklendi)", "“Ben bunu yapabilirim” inancı", "“Bırakabilirim.”"],
], [4.0, 5.5, 6.0])
not_kutusu("Bağlantı: Hocanın 53. slaytındaki 5R'de yer alan Risks (riskler), Rewards (ödüller) ve Roadblocks (engeller), bu modelin algılanan ciddiyet, yarar ve engel bileşenlerine karşılık gelmektedir.")

# ---------------- 2 ----------------
par("2. Sağlık Kontrol Odağı", style="Heading 1")
par("**Kim, ne zaman:** 1970'lerde Wallston ve arkadaşları, Rotter'ın “kontrol odağı” kavramını sağlığa uyarlamıştır.")
par("**Ana fikir:** Kişinin “Sağlığımı kim ya da ne belirler?” sorusuna verdiği cevap, sağlık davranışlarını etkiler.")
etiket("Üç boyutu")
madde("**İç kontrol odağı:** “Sağlığım benim elimde.” Sorumluluk alır, bırakmaya daha istekli olma eğilimindedir.")
madde("**Dış kontrol, güçlü başkaları:** “Sağlığımı doktorum ya da ailem belirler.” Hekimin yönlendirmesi ve yakın takip daha etkilidir.")
madde("**Dış kontrol, şans:** “Kanser olacaksam olurum, kader.” Değişime en dirençli gruptur.")
etiket("Pratikte")
par("Hekim, hastanın hangi tipte olduğuna göre yaklaşır. İç kontrollü hastaya sorumluluğu verir. Dış kontrollü hastada daha sık görüşme planlar ve hastayı yönlendirir.")

# ---------------- 3 ----------------
par("3. Yeniliklerin Yayılımı Modeli", style="Heading 1")
par("**Kim, ne zaman:** Everett Rogers, 1962.")
par("**Ana fikir:** Yeni bir fikir, davranış veya ürün toplumda herkese aynı anda değil, gruplar halinde ve zamanla yayılır. Sigarasız yaşam da toplum için bir “yenilik” olarak düşünülebilir.")
etiket("Benimseyen gruplar")
tablo(["Grup", "Oran"], [
    ["Yenilikçiler", "%2,5"],
    ["Erken benimseyenler (kanaat önderleri)", "%13,5"],
    ["Erken çoğunluk", "%34"],
    ["Geç çoğunluk", "%34"],
    ["Geride kalanlar", "%16"],
], [8.0, 3.0])
etiket("Yeniliğin yayılmasını kolaylaştıran özellikler")
for t in ["Göreceli avantajı olması", "Kişinin yaşamına uygun olması", "Basit olması", "Denenebilir olması", "Sonuçlarının görülebilir olması"]:
    madde(t)
etiket("Sigara örneği")
par("Hekimler, öğretmenler ve din görevlileri gibi kanaat önderleri sigarayı bıraktığında, çevrelerindeki çoğunluk da onları izler.")
not_kutusu("Bağlantı: Hocanın 47. slaytındaki “rol modeli olma” vurgusu bu modelin mantığına uymaktadır.")

# ---------------- 4 ----------------
par("4. Pender'in Sağlığı Geliştirme Modeli", style="Heading 1")
par("**Kim, ne zaman:** Hemşire Nola Pender, 1982 (1996'da yeniden düzenlenmiştir). Hemşirelik kökenli bir modeldir.")
par("**Ana fikir:** Sağlık İnanç Modeli'nden farkı, korkuya veya hastalık tehdidine değil, olumlu bir hedefe, yani daha sağlıklı yaşama isteğine dayanmasıdır. Kişi “hastalanmamak için” değil, “daha iyi yaşamak için” davranışını değiştirir.")
etiket("Üç ana bölümü")
numara(1, "**Bireysel özellikler ve deneyimler**")
madde("Önceki ilgili davranışlar (örneğin önceki bırakma denemeleri)", 1)
madde("Yaş, eğitim gibi kişisel faktörler", 1)
numara(2, "**Davranışa özgü düşünce ve duygular**")
madde("Algılanan yarar ve engel", 1)
madde("Öz-yeterlilik: “Yapabilirim.”", 1)
madde("Kişilerarası etkiler: aile, arkadaş, sağlık çalışanı", 1)
madde("Durumsal etkiler: ortam", 1)
numara(3, "**Davranış sonucu**")
madde("Eylem planına bağlılık (bırakma tarihi belirlemek)", 1)
madde("O anki rekabet eden istekler (sigara isteği)", 1)
madde("Sonuçta sağlığı geliştirici davranış: sigarayı bırakmak", 1)
etiket("Sunulan makaleyle örtüşen noktalar")
madde("**Kişilerarası etki:** Makalede hanede sigara içen olmasının bırakma başarısını düşürdüğü bulunmuştur.")
madde("**İçsel ve olumlu motivasyon:** Makalede yakınması olmayanların daha başarılı olması, “anlık tepkisel korkudan ziyade uzun vadeli sağlığı koruma amacıyla içsel motivasyon” ile açıklanmaktadır.")
not_kutusu("Bağlantı: Makale bu bulguları Pender'in modeliyle ilişkilendirmemektedir; benzerlik tarafımdan kurulmuştur.")

# ---------------- 5 ----------------
par("5. Transteorik Model", style="Heading 1")
par("**Kim, ne zaman:** Prochaska ve DiClemente, 1983. “Değişim aşamaları modeli” olarak da bilinir.")
par("**Ana fikir:** Davranış değişikliği bir anda olmaz; kişi belirli aşamalardan geçer. Hekimin görevi hastanın hangi aşamada olduğunu belirleyip ona uygun müdahaleyi yapmaktır.")
etiket("Hocanın slaytındaki basamaklar (Slayt 50)")
numara(1, "**Kişide davranış değişikliği niyetinin olmadığı, farkındalık öncesi; düşünme öncesi aşama:** Hasta bırakmayı düşünmüyor. Bu hastaya hocanın 53. slaytındaki 5R yaklaşımı uygulanır (“sigarayı bırakmayı düşünmeyen kişilere uygulanacak girişimler”).")
numara(2, "**Niyet aşaması:** Bırakmayı düşünüyor ama kararsız; “bırakmalıyım ama…” diyor.")
numara(3, "**Hazırlık aşaması:** Bırakmaya karar vermiş, bırakma tarihi belirliyor. Tedavi bu aşamada planlanır.")
numara(4, "**Eylem aşaması:** Sigarayı bırakmış. Craving ile mücadele için hocanın 54. slaytındaki 4D (geciktir, derin nefes al, su iç, dikkati dağıt) burada önemlidir.")
numara(5, "**Sürdürme aşaması:** Bırakmayı uzun süredir koruyor.")
numara(6, "**Nüks aşaması:** Yeniden sigaraya başlamış.")
madde("Slayt 73: Tek sigara veya birkaç puff “kayma” (slip), birkaç gün birer sigara “hata” (laps) olarak tanımlanır; ikisi de nüksün ön göstergesidir.", 1)
madde("Slayt 74: Nüks en sık ilk haftada görülür; ilk yılda oran %60–90'dır.", 1)
etiket("Genel literatürden ek bilgi (hocanın slaytında yazmıyor)")
madde("Aşamalar zamana göre tanımlanır: niyet aşamasında “6 ay içinde bırakmayı düşünüyor”, hazırlıkta “1 ay içinde”, eylemde “6 aydan kısa süredir bırakmış”, sürdürmede “6 aydan uzun süredir bırakmış”.")
madde("Nüks eden kişi başa dönmez, döngüye yeniden girer; bu yüzden model bir döngü olarak çizilir.")
etiket("Hocanın dersindeki bağlantılar")
madde("**5A'daki “Assess” (Slayt 49):** Hastanın hangi aşamada olduğu bu adımda değerlendirilir.")
madde("**Ders başındaki soru (Slayt 5):** “Hangisi Transteorik modelin basamaklarından biri değildir?” Cevap: “Vazgeçme aşaması”. Böyle bir aşama yoktur.")
madde("**Kaynak 4:** Fidanci I, Ozturk O, Unal M. Transtheoretic Model in smoking cessation. Onur hocanın kendi yayınıdır; soru gelme olasılığı en yüksek konu budur.")

# ---------------- Özet ----------------
par("Akılda tutmak için", style="Heading 1")
tablo(["Model", "Anahtar fikir"], [
    ["Sağlık İnanç Modeli", "Korku ve tehdit"],
    ["Sağlık Kontrol Odağı", "Kontrol kimde?"],
    ["Yeniliklerin Yayılımı", "Toplumda yayılma, rol model"],
    ["Pender", "Olumlu hedef, sağlıklı yaşam isteği"],
    ["Transteorik Model", "Aşamalar; hastanın aşamasına göre müdahale"],
], [6.0, 9.5])

d.save(CIKIS)
print("ok")
